"""Read uploaded files into plain text.

The file kind comes from the extension and, when they disagree, from the
file header. A PDF with a real text layer is read locally. A PDF that only
wraps a scan is treated as that scan: the embedded photo is sent to the
vision callback at its own size, not redrawn as a small page image. Legacy
.doc is rejected with a clear message instead of being treated as empty text.
"""

import os
import tempfile
import zipfile

_EXT_KIND = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "doc",
    ".jpg": "jpeg",
    ".jpeg": "jpeg",
    ".png": "png",
    ".webp": "webp",
    ".gif": "gif",
    ".tif": "tiff",
    ".tiff": "tiff",
    ".bmp": "bmp",
    ".txt": "text",
    ".csv": "csv",
}

IMAGE_KINDS = {"jpeg", "png", "webp", "gif", "tiff", "bmp"}
_DIRECT_IMAGE_KINDS = {"jpeg", "png", "webp", "gif"}
_MIN_TEXT_CHARS = 40
# A wrapped scan sometimes carries a short fake text layer (a title, a page
# number). A real invoice text layer is much longer than this.
_MIN_REAL_PDF_CHARS = 400
_SCAN_PAGE_COVER = 0.5
_FALLBACK_DPI = 200
_DEFAULT_PDF_PAGES = 8


def detect_kind(file_path):
    """Return a short kind name such as pdf, docx, jpeg, or unknown."""
    ext_kind = _EXT_KIND.get(os.path.splitext(file_path)[1].lower())
    magic = _magic_kind(file_path)
    if magic and ext_kind and magic != ext_kind:
        return magic
    return ext_kind or magic or "unknown"


def extract_file_text(file_path, ocr_image, max_pdf_pages=_DEFAULT_PDF_PAGES):
    """Return (text, source_format).

    ocr_image(path) transcribes one image file and returns text.
    """
    if not os.path.isfile(file_path):
        raise ValueError(f"File not found: {os.path.basename(file_path)}")

    kind = detect_kind(file_path)
    if kind == "pdf":
        text = _pdf_text_layer(file_path)
        if _pdf_text_is_real(file_path, text):
            return text, "pdf-text"
        scanned = _ocr_pdf(file_path, ocr_image, max_pdf_pages)
        if not _usable(scanned):
            raise ValueError(
                "This PDF has no text layer, and the vision model returned no text. "
                "Enable a vision-capable provider (OpenRouter, Gemini, Claude, or OpenAI) "
                "with an image model, then upload the file again."
            )
        return scanned, "pdf-scan"
    if kind == "docx":
        text = _docx_text(file_path)
        if not _usable(text):
            raise ValueError("The DOCX file has no readable text.")
        return text, "docx"
    if kind in IMAGE_KINDS:
        text = _ocr_image(file_path, kind, ocr_image)
        if not _usable(text):
            raise ValueError(
                "The image was sent for OCR, but no text came back. "
                "Check that the provider model accepts images."
            )
        return text, kind
    if kind in {"text", "csv"}:
        text = _plain_text(file_path)
        if not _usable(text):
            raise ValueError("The text file is empty.")
        return text, kind
    if kind == "doc":
        raise ValueError(
            "Legacy Word .doc files are not read. Save the file as DOCX or PDF and upload it again."
        )
    raise ValueError(
        f"Unsupported file type ({kind}). Supported: PDF, DOCX, TXT, CSV, "
        "and images (JPG, PNG, WEBP, GIF, TIFF, BMP)."
    )


def _usable(text):
    return len("".join((text or "").split())) >= _MIN_TEXT_CHARS


def _magic_kind(file_path):
    try:
        with open(file_path, "rb") as handle:
            header = handle.read(16)
    except OSError:
        return None
    if header.startswith(b"%PDF"):
        return "pdf"
    if header.startswith(b"\xff\xd8"):
        return "jpeg"
    if header.startswith(b"\x89PNG"):
        return "png"
    if header.startswith(b"GIF8"):
        return "gif"
    if header.startswith(b"BM"):
        return "bmp"
    if header.startswith(b"RIFF") and _looks_like_webp(file_path):
        return "webp"
    if header.startswith(b"\xd0\xcf\x11\xe0"):
        return "doc"
    if header.startswith(b"PK") and _zip_has_docx(file_path):
        return "docx"
    if header.startswith(b"II*\x00") or header.startswith(b"MM\x00*"):
        return "tiff"
    return None


def _looks_like_webp(file_path):
    try:
        with open(file_path, "rb") as handle:
            chunk = handle.read(16)
    except OSError:
        return False
    return len(chunk) >= 12 and chunk[8:12] == b"WEBP"


def _zip_has_docx(file_path):
    try:
        with zipfile.ZipFile(file_path) as archive:
            names = set(archive.namelist())
    except (OSError, zipfile.BadZipFile):
        return False
    return "word/document.xml" in names


def _pdf_text_layer(file_path):
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception:
        return ""


def _docx_text(file_path):
    from docx import Document
    document = Document(file_path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _plain_text(file_path):
    with open(file_path, "rb") as handle:
        raw = handle.read()
    for encoding in ("utf-8-sig", "utf-8", "cp1256", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _ocr_image(file_path, kind, ocr_image):
    if kind in _DIRECT_IMAGE_KINDS:
        return ocr_image(file_path) or ""
    jpeg_path = _render_jpeg(file_path)
    try:
        return ocr_image(jpeg_path) or ""
    finally:
        if os.path.exists(jpeg_path):
            os.remove(jpeg_path)


def _import_pymupdf():
    try:
        import pymupdf
    except ImportError as exc:
        raise ValueError(
            "pymupdf is required to read a scanned PDF or a TIFF/BMP image."
        ) from exc
    return pymupdf


def _pdf_text_is_real(file_path, text):
    """A short text layer on top of a full-page photo is not the invoice."""
    compact = len("".join((text or "").split()))
    if compact < _MIN_TEXT_CHARS:
        return False
    if compact >= _MIN_REAL_PDF_CHARS:
        return True
    return not _pdf_is_wrapped_scan(file_path)


def _pdf_is_wrapped_scan(file_path):
    try:
        pymupdf = _import_pymupdf()
    except ValueError:
        return False
    document = pymupdf.open(file_path)
    try:
        if document.page_count < 1:
            return False
        limit = min(document.page_count, _DEFAULT_PDF_PAGES)
        wrapped = 0
        for index in range(limit):
            if _dominant_scan_xref(document[index]):
                wrapped += 1
        return wrapped > 0 and wrapped == limit
    except Exception:
        return False
    finally:
        document.close()


def _ocr_pdf(file_path, ocr_image, max_pages):
    pymupdf = _import_pymupdf()
    document = pymupdf.open(file_path)
    parts = []
    try:
        total = document.page_count
        limit = min(total, max_pages)
        for index in range(limit):
            page = document[index]
            image_path = _scan_image_file(document, page)
            try:
                text = (ocr_image(image_path) or "").strip()
            finally:
                if os.path.exists(image_path):
                    os.remove(image_path)
            if text:
                parts.append(f"--- page {index + 1} ---\n{text}")
        if total > limit:
            parts.append(f"[Stopped after {limit} of {total} pages.]")
    finally:
        document.close()
    return "\n\n".join(parts)


def _scan_image_file(document, page):
    """Prefer the photo stored inside the page. Redraw only when it is missing."""
    xref = _dominant_scan_xref(page)
    if xref:
        extracted = _write_embedded_image(document, xref)
        if extracted:
            return extracted
    return _pixmap_to_jpeg(page, dpi=_FALLBACK_DPI)


def _dominant_scan_xref(page):
    """Xref of the picture that fills the page, or None for a normal text page."""
    page_area = abs(page.rect.width * page.rect.height) or 1
    best_pixels = 0
    best_xref = None
    infos = []
    try:
        infos = page.get_image_info(xrefs=True) or []
    except Exception:
        infos = []
    for info in infos:
        xref = info.get("xref") or 0
        if xref <= 0:
            continue
        bbox = info.get("bbox")
        cover = _box_area(bbox) / page_area
        if cover < _SCAN_PAGE_COVER:
            continue
        pixels = int(info.get("width") or 0) * int(info.get("height") or 0)
        if pixels >= best_pixels:
            best_pixels = pixels
            best_xref = xref
    if best_xref:
        return best_xref
    try:
        images = page.get_images(full=True) or []
    except Exception:
        return None
    for item in images:
        xref = item[0]
        width = item[2] or 0
        height = item[3] or 0
        pixels = width * height
        if min(width, height) >= 800 and pixels > best_pixels:
            best_pixels = pixels
            best_xref = xref
    return best_xref


def _box_area(bbox):
    if bbox is None:
        return 0
    if hasattr(bbox, "width"):
        return abs(bbox.width * bbox.height)
    try:
        return abs((bbox[2] - bbox[0]) * (bbox[3] - bbox[1]))
    except (TypeError, IndexError):
        return 0


def _write_embedded_image(document, xref):
    try:
        extracted = document.extract_image(xref)
    except Exception:
        return None
    if not extracted or not extracted.get("image"):
        return None
    ext = (extracted.get("ext") or "").lower()
    if ext == "jpeg":
        ext = "jpg"
    if ext not in {"jpg", "png", "webp", "gif"}:
        return None
    handle = tempfile.NamedTemporaryFile(suffix="." + ext, delete=False)
    try:
        handle.write(extracted["image"])
    finally:
        handle.close()
    return handle.name


def _render_jpeg(file_path):
    pymupdf = _import_pymupdf()
    document = pymupdf.open(file_path)
    try:
        if document.page_count < 1:
            raise ValueError("The image file has no pages to read.")
        return _pixmap_to_jpeg(document[0])
    finally:
        document.close()


def _pixmap_to_jpeg(page, dpi=_FALLBACK_DPI):
    pixmap = page.get_pixmap(dpi=dpi)
    handle = tempfile.NamedTemporaryFile(suffix=".jpg", delete=False)
    handle.close()
    pixmap.save(handle.name)
    return handle.name
